"""
Claude API service — handles all AI operations:
- Extracting and simplifying content from uploaded files
- Generating syllabi and lesson plans
- Generating quiz questions
- Evaluating knowledge level from quiz results
"""
import json
import re
import fitz  # PyMuPDF
from pptx import Presentation
import anthropic

from app.config import settings

_client = anthropic.Anthropic(api_key=settings.anthropic_api_key)
_MODEL = "claude-haiku-4-5-20251001"  # fast + cheap for hackathon; swap to claude-sonnet-4-6 for quality


def _ask(prompt: str, max_tokens: int = 4096) -> str:
    message = _client.messages.create(
        model=_MODEL,
        max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}],
    )
    return message.content[0].text


# ---------------------------------------------------------------------------
# File text extraction
# ---------------------------------------------------------------------------

def extract_text_from_pdf(file_path: str) -> str:
    doc = fitz.open(file_path)
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n\n".join(pages)


def extract_text_from_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        slide_text = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_text.append(text)
        slides.append("\n".join(slide_text))
    return "\n\n".join(slides)


def extract_text(file_path: str, file_type: str) -> str:
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type in ("pptx", "ppt"):
        return extract_text_from_pptx(file_path)
    raise ValueError(f"Unsupported file type: {file_type}")


# ---------------------------------------------------------------------------
# JSON parsing helper
# ---------------------------------------------------------------------------

def _parse_json_response(text: str) -> dict | list:
    """Strip markdown fences and parse JSON from Claude's response."""
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return json.loads(text)


# ---------------------------------------------------------------------------
# Syllabus + lesson plan generation
# ---------------------------------------------------------------------------

SYLLABUS_PROMPT = """
You are an expert educator. Analyse the following study material and produce a structured syllabus.

Return ONLY valid JSON (no markdown, no explanation) in this exact format:
{{
  "summary": {{
    "subject": "<subject title>",
    "overview": "<2-3 sentence overview of the material>",
    "total_topics": <number>,
    "estimated_total_hours": <number>
  }},
  "topics": [
    {{
      "topic_number": 1,
      "title": "<topic title>",
      "summary": "<plain-English explanation, 3-5 sentences, suitable for a student new to the topic>",
      "key_concepts": ["<concept 1>", "<concept 2>", ...],
      "estimated_minutes": <number>
    }}
  ],
  "lesson_plan": [
    {{
      "day": 1,
      "date_label": "Day 1",
      "topics": ["<topic title>"],
      "duration_minutes": <number>
    }}
  ]
}}

Rules:
- Simplify complex academic language into clear, accessible explanations
- Each topic should be self-contained and learnable in one session
- The lesson plan should spread topics across days (max 90 minutes per day)
- key_concepts should be 3-7 bullet points per topic

Study material:
{content}
"""


async def generate_syllabus(file_path: str, file_type: str) -> dict:
    content = extract_text(file_path, file_type)

    # Truncate to avoid token limits (~50k chars ≈ ~12k tokens)
    if len(content) > 50000:
        content = content[:50000] + "\n\n[Content truncated for processing]"

    prompt = SYLLABUS_PROMPT.format(content=content)
    response = _ask(prompt, max_tokens=8192)
    return _parse_json_response(response)


# ---------------------------------------------------------------------------
# Quiz generation
# ---------------------------------------------------------------------------

ASSESSMENT_QUIZ_PROMPT = """
You are a quiz designer. Based on the syllabus below, create a knowledge assessment quiz
to determine the student's existing knowledge level.

Return ONLY valid JSON (no markdown) in this exact format:
{{
  "questions": [
    {{
      "question": "<question text>",
      "options": ["<option A>", "<option B>", "<option C>", "<option D>"],
      "correct_index": <0-3>,
      "explanation": "<brief explanation of why the answer is correct>"
    }}
  ]
}}

Rules:
- Create exactly {num_questions} multiple-choice questions
- Mix easy, medium, and hard questions
- Cover all major topics from the syllabus
- Each question must have exactly 4 options
- correct_index is 0-based (0=A, 1=B, 2=C, 3=D)

Syllabus:
{syllabus_json}
"""

LESSON_QUIZ_PROMPT = """
You are a quiz designer. Based on the topic below from a syllabus, create a quiz to test
understanding of that specific topic.

Return ONLY valid JSON (no markdown) in this exact format:
{{
  "questions": [
    {{
      "question": "<question text>",
      "options": ["<option A>", "<option B>", "<option C>", "<option D>"],
      "correct_index": <0-3>,
      "explanation": "<brief explanation of why the answer is correct>"
    }}
  ]
}}

Rules:
- Create exactly {num_questions} multiple-choice questions
- Focus specifically on the topic provided
- Vary difficulty from recall to application
- Each question must have exactly 4 options

Topic: {topic_title}
Topic content:
{topic_json}
"""


async def generate_assessment_quiz(syllabus_data: dict, num_questions: int = 10) -> list[dict]:
    prompt = ASSESSMENT_QUIZ_PROMPT.format(
        num_questions=num_questions,
        syllabus_json=json.dumps(syllabus_data, indent=2),
    )
    response = _ask(prompt)
    data = _parse_json_response(response)
    return data["questions"]


async def generate_lesson_quiz(topic: dict, num_questions: int = 5) -> list[dict]:
    prompt = LESSON_QUIZ_PROMPT.format(
        num_questions=num_questions,
        topic_title=topic["title"],
        topic_json=json.dumps(topic, indent=2),
    )
    response = _ask(prompt)
    data = _parse_json_response(response)
    return data["questions"]


# ---------------------------------------------------------------------------
# Level assessment from quiz results
# ---------------------------------------------------------------------------

LEVEL_ASSESSMENT_PROMPT = """
A student just completed a knowledge assessment quiz with a score of {score_percent}%.
They got {correct}/{total} questions correct.

Based on this performance, provide a brief personalised learning assessment.

Return ONLY valid JSON (no markdown):
{{
  "level": "<Beginner|Intermediate|Advanced>",
  "assessment": "<2-3 sentence personalised assessment of their knowledge level>",
  "recommendation": "<1-2 sentence recommendation for how to approach the study plan>"
}}
"""


async def assess_knowledge_level(score: float, correct: int, total: int) -> dict:
    prompt = LEVEL_ASSESSMENT_PROMPT.format(
        score_percent=round(score * 100),
        correct=correct,
        total=total,
    )
    response = _ask(prompt, max_tokens=512)
    return _parse_json_response(response)


# ---------------------------------------------------------------------------
# Lesson content generation
# ---------------------------------------------------------------------------

LESSON_CONTENT_PROMPT = """
You are an expert teacher creating a detailed lesson page for a student.

Topic: {topic_title}
Topic summary: {topic_summary}
Key concepts: {key_concepts}
Subject context: {subject}

Create a multi-section lesson that thoroughly explains this topic to a beginner-to-intermediate student.

Return ONLY valid JSON (no markdown) in this exact format:
{{
  "sections": [
    {{
      "type": "intro",
      "title": "Introduction",
      "content": "<engaging 3-4 sentence introduction that hooks the student and explains why this topic matters>"
    }},
    {{
      "type": "explanation",
      "title": "<descriptive title for core explanation>",
      "content": "<thorough explanation of the main concepts, 5-8 sentences, use plain language>"
    }},
    {{
      "type": "explanation",
      "title": "<title for second key idea>",
      "content": "<explanation of another key aspect, 4-6 sentences>"
    }},
    {{
      "type": "example",
      "title": "Real-World Example",
      "content": "<concrete, relatable example that illustrates the concept, 4-5 sentences>"
    }},
    {{
      "type": "tips",
      "title": "Key Takeaways",
      "content": "<3-5 bullet points separated by newlines, starting each with • , summarising the most important points>"
    }}
  ]
}}

Rules:
- Write in a friendly, encouraging tone
- Avoid jargon — explain terms when you use them
- Keep each section focused and concise
- The example should be something a student can easily picture
"""

LESSON_QA_SYSTEM = """You are a helpful tutor assisting a student who is studying the following topic.

Topic: {topic_title}
Topic summary: {topic_summary}
Key concepts: {key_concepts}

Answer questions clearly and concisely. If the student asks something off-topic, gently redirect them to the topic at hand. Keep answers under 150 words unless the question requires detail."""


async def generate_lesson(topic: dict, subject: str = "") -> list[dict]:
    """Generate multi-section lesson content for a topic."""
    prompt = LESSON_CONTENT_PROMPT.format(
        topic_title=topic["title"],
        topic_summary=topic.get("summary", ""),
        key_concepts=", ".join(topic.get("key_concepts", [])),
        subject=subject,
    )
    response = _ask(prompt, max_tokens=4096)
    data = _parse_json_response(response)
    return data["sections"]


async def answer_question(
    topic: dict,
    question: str,
    history: list[dict],
) -> str:
    """Answer a student question about a topic, with conversation history."""
    system = LESSON_QA_SYSTEM.format(
        topic_title=topic["title"],
        topic_summary=topic.get("summary", ""),
        key_concepts=", ".join(topic.get("key_concepts", [])),
    )

    messages = []
    for msg in history:
        messages.append({"role": msg["role"], "content": msg["content"]})
    messages.append({"role": "user", "content": question})

    response = _client.messages.create(
        model=_MODEL,
        max_tokens=512,
        system=system,
        messages=messages,
    )
    return response.content[0].text
